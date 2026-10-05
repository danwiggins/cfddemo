"""Bounded orchestration for Loyfer fragment-UXM cell-origin regeneration.

The public pipeline accepts either a validated, normalized modkit extract TSV
or an aligned modBAM that can be extracted by a verified modkit executable.
It joins the supplied Loyfer marker registry, region BED, and U atlas; performs
fragment UXM classification; fits count-weighted NNLS; and publishes an atomic,
aggregate-only JSON result bundle.

Loyfer Supplementary Table S8 provides the method-matched observed plasma
cohort used for chart context.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import platform
import re
import shutil
import subprocess
import sys
import tempfile
import threading
from bisect import bisect_right
from collections import defaultdict
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Literal

import numpy as np
from pydantic import Field, TypeAdapter, ValidationError, model_validator

from evidence_inspector.cell_origin_inputs import (
    AtlasUColumns,
    CellOriginInputError,
    FractionUnit,
    MarkerBedColumns,
    ModkitExtractColumns,
    load_loyfer_atlas_u_matrix,
    load_marker_bed,
    load_modkit_extract_calls,
)
from evidence_inspector.cell_origin_models import (
    AtlasUMatrix,
    BootstrapInformationStatus,
    BootstrapResultV2,
    CellOriginProvenance,
    CellOriginResult,
    DeconvolutionOutputV2,
    DigestArtifact,
    GenomicMarker,
    Identifier,
    LOYFER_UXM_METHOD,
    MarkerCountRow,
    NnlsRowScale,
    SoftwareVersion,
    StrictModel,
    UxmThresholds,
    ValidationCheck,
    ValidationRecord,
    ValidationReport,
    VerificationLevel,
)
from evidence_inspector.cell_origin_prefilter import (
    DEFAULT_MIN_MAPQ,
    PrefilterCounts,
    prefilter_summary,
    read_bed_regions,
    write_prefiltered_bam,
)
from evidence_inspector.deconvolution import (
    bootstrap_uxm_v2,
    deconvolve_uxm_v2,
)
from evidence_inspector.uxm import (
    UxmClassificationResult,
    UxmDiagnostics,
    UxmStopReason,
    aggregate_marker_counts,
    classify_uxm_calls,
)
from traceback_runner.toolchain import (
    PinnedTool,
    ToolProblem,
    exec_pinned,
    kill_process_group,
    resolve_modkit,
)

PIPELINE_VERSION = "1.0.0"
SCHEMA_VERSION = "cell-origin-pipeline.v1"
DEFAULT_MAXIMUM_CALLS = 1_000_000
DEFAULT_MAXIMUM_GROUPS = 100_000
DEFAULT_MAXIMUM_CPGS_PER_GROUP = 10_000
DEFAULT_BOOTSTRAP_REPLICATES = 200
DEFAULT_RANDOM_SEED = 7
DEFAULT_TOP_COMPOSITION_ROWS = 12
# modkit's global ``--filter-threshold``, locked: never modkit's per-run
# estimate.  It is the C threshold modkit 0.6.4 estimated for the reproduced
# 2026-10-01 run (qual 233, i.e. (233 + 0.5) / 256), written exactly.  With it,
# the C calls modkit emits are the same as with the estimate; only non-C rows,
# which the call loader drops, can change.  A method parameter, not a result.
DEFAULT_MODKIT_FILTER_THRESHOLD = 0.912109375
MODKIT_TIMEOUT_SECONDS = 1800
# How long a stream that ended early waits for modkit's own exit status.
MODKIT_EXIT_GRACE_SECONDS = 5
MODKIT_WORK_DIRECTORY = "modkit-work"
PALETTE = (
    "#0F766E",
    "#2563EB",
    "#7C3AED",
    "#DB2777",
    "#EA580C",
    "#CA8A04",
    "#16A34A",
    "#0891B2",
    "#4F46E5",
    "#9333EA",
    "#E11D48",
    "#475569",
)

NORMALIZED_MODKIT_COLUMNS = ModkitExtractColumns(
    fragment_id="read_id",
    chromosome="chrom",
    position0="ref_position",
    strand="mod_strand",
    modified_primary_base="modified_primary_base",
    call_code="call_code",
    modified_probability="modified_probability",
    fail="fail",
)

# Row cap of every Loyfer resource file and of the derived atlas and marker tables.
LOYFER_MAX_ROWS = 20_000
MARKER_METADATA_HEADER = (
    "#chr",
    "start",
    "end",
    "startCpG",
    "endCpG",
    "target",
    "region",
    "lenCpG",
    "bp",
    "tg_mean",
    "bg_mean",
    "delta_means",
    "delta_quants",
    "delta_maxmin",
    "ttest",
    "direction",
)
ATLAS_METADATA_HEADER = (
    "chr",
    "start",
    "end",
    "startCpG",
    "endCpG",
    "target",
    "name",
    "direction",
)
NATIVE_MODKIT_REQUIRED_COLUMNS = frozenset(
    {
        "read_id",
        "ref_position",
        "chrom",
        "mod_strand",
        "modified_primary_base",
        "fail",
        "call_code",
    }
)


class CellOriginPipelineError(RuntimeError):
    """Sanitized, actionable pipeline failure."""


class CallCapExceeded(CellOriginPipelineError):
    """The modkit extract passed the locked call cap; the run is refused."""


class NoMarkerAlignments(CellOriginPipelineError):
    """No alignment passed the pre-filter and overlapped a marker region."""


class ModkitTimedOut(CellOriginPipelineError):
    """modkit did not finish within its timeout; its process group was killed."""


class ModkitKilled(CellOriginPipelineError):
    """modkit was stopped by a signal (for example the OS under memory pressure)."""


class CommandStep(StrictModel):
    """One directly executable argv vector; never a shell command string."""

    name: str = Field(min_length=1, max_length=64)
    argv: tuple[str, ...] = Field(min_length=1)
    purpose: str = Field(min_length=1, max_length=512)


class AlignmentCommandPlan(StrictModel):
    """Auditable minimap2 plan that carries MM/ML/MN via FASTQ comments."""

    steps: tuple[CommandStep, ...] = Field(min_length=1)
    preserved_tags: tuple[Literal["MM", "ML", "MN"], ...] = ("MM", "ML", "MN")
    postconditions: tuple[str, ...] = Field(min_length=1)


class PreflightItem(StrictModel):
    item_id: str = Field(
        min_length=1,
        max_length=128,
        pattern=r"^[A-Za-z0-9][A-Za-z0-9_.:-]*$",
    )
    ready: bool
    blocking: bool
    version: str | None = Field(default=None, max_length=128)
    detail: str = Field(min_length=1, max_length=512)


class PreflightReport(StrictModel):
    items: tuple[PreflightItem, ...] = Field(min_length=1)

    @property
    def ready(self) -> bool:
        return not any(item.blocking and not item.ready for item in self.items)

    @property
    def blockers(self) -> tuple[str, ...]:
        return tuple(
            item.detail
            for item in self.items
            if item.blocking and not item.ready
        )


class CompositionChartRow(StrictModel):
    cell_type_id: str
    label: str
    rank: int = Field(ge=1)
    fraction: float = Field(ge=0.0, le=1.0)
    percent: float = Field(ge=0.0, le=100.0)
    lower_fraction: float | None = Field(default=None, ge=0.0, le=1.0)
    upper_fraction: float | None = Field(default=None, ge=0.0, le=1.0)
    lower_percent: float | None = Field(default=None, ge=0.0, le=100.0)
    upper_percent: float | None = Field(default=None, ge=0.0, le=100.0)
    uncertainty_available: bool = True
    uncertainty_status: BootstrapInformationStatus = (
        BootstrapInformationStatus.AVAILABLE
    )
    color: str = Field(pattern=r"^#[0-9A-F]{6}$")
    show_by_default: bool

    @model_validator(mode="after")
    def validate_interval(self) -> CompositionChartRow:
        values = (
            self.lower_fraction,
            self.upper_fraction,
            self.lower_percent,
            self.upper_percent,
        )
        has_interval = all(value is not None for value in values)
        if any(value is not None for value in values) and not has_interval:
            raise ValueError("composition uncertainty bounds must be all-or-none")
        if self.uncertainty_available != has_interval:
            raise ValueError("composition uncertainty availability is inconsistent")
        if self.uncertainty_available:
            if self.uncertainty_status != BootstrapInformationStatus.AVAILABLE:
                raise ValueError(
                    "available composition uncertainty needs available status"
                )
            assert self.lower_fraction is not None
            assert self.upper_fraction is not None
            if not self.lower_fraction <= self.fraction <= self.upper_fraction:
                raise ValueError("composition uncertainty must contain the estimate")
        elif self.uncertainty_status == BootstrapInformationStatus.AVAILABLE:
            raise ValueError("unavailable composition uncertainty cannot be available")
        return self


class HealthyRangeChartRow(StrictModel):
    """Method-matched observed plasma range from Loyfer Supplementary Table S8."""

    cell_type_id: str
    label: str
    sample_fraction: float = Field(ge=0.0, le=1.0)
    sample_percent: float = Field(ge=0.0, le=100.0)
    sample_lower_percent: float | None = Field(default=None, ge=0.0, le=100.0)
    sample_upper_percent: float | None = Field(default=None, ge=0.0, le=100.0)
    healthy_min_percent: float = Field(ge=0.0, le=100.0)
    healthy_q1_percent: float = Field(ge=0.0, le=100.0)
    healthy_median_percent: float = Field(ge=0.0, le=100.0)
    healthy_q3_percent: float = Field(ge=0.0, le=100.0)
    healthy_max_percent: float = Field(ge=0.0, le=100.0)
    healthy_mean_percent: float = Field(ge=0.0, le=100.0)
    comparison_valid: Literal[True] = True
    classification: Literal["below", "within", "above"]
    uncertainty_available: bool
    reference_label: Literal["Loyfer Table S8 plasma donors (n=23)"] = (
        "Loyfer Table S8 plasma donors (n=23)"
    )

    @model_validator(mode="after")
    def validate_ranges(self) -> HealthyRangeChartRow:
        if not (
            self.healthy_min_percent
            <= self.healthy_q1_percent
            <= self.healthy_median_percent
            <= self.healthy_q3_percent
            <= self.healthy_max_percent
        ):
            raise ValueError("healthy chart statistics must be ordered")
        if self.uncertainty_available != (
            self.sample_lower_percent is not None
            and self.sample_upper_percent is not None
        ):
            raise ValueError("uncertainty availability does not match interval")
        return self


class ChartData(StrictModel):
    composition_rows: tuple[CompositionChartRow, ...] = Field(min_length=1)
    healthy_context_rows: tuple[HealthyRangeChartRow, ...] = ()
    composition_title: str
    range_title: str | None = None
    x_axis_title: Literal["Estimated contribution (%)"] = (
        "Estimated contribution (%)"
    )
    method_warning: str | None = None


class ResourceSummary(StrictModel):
    registered_markers: int = Field(ge=1)
    usable_markers: int = Field(ge=1)
    excluded_incomplete_atlas_markers: int = Field(ge=0)
    collapsed_duplicate_regions: int = Field(ge=0)
    cell_type_count: int = Field(ge=1)


class CellOriginResultBundle(StrictModel):
    schema_version: Literal["cell-origin-pipeline.v1"] = SCHEMA_VERSION
    result: CellOriginResult
    charts: ChartData
    resources: ResourceSummary
    notices: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class PipelineConfig:
    marker_bed: Path
    marker_metadata: Path
    atlas_u_matrix: Path
    output_path: Path
    extract_tsv: Path | None = None
    aligned_modbam: Path | None = None
    healthy_table: Path | None = None
    healthy_sheet: str | None = None
    healthy_fraction_unit: FractionUnit = FractionUnit.FRACTION
    atlas_id: str = "atlas.loyfer-u250-l4.hg38"
    atlas_source_id: str = "source.loyfer-u250"
    healthy_source_id: str = "source.loyfer-table-s8"
    maximum_calls: int = DEFAULT_MAXIMUM_CALLS
    maximum_groups: int = DEFAULT_MAXIMUM_GROUPS
    maximum_cpgs_per_group: int = DEFAULT_MAXIMUM_CPGS_PER_GROUP
    bootstrap_replicates: int = DEFAULT_BOOTSTRAP_REPLICATES
    random_seed: int = DEFAULT_RANDOM_SEED
    # Required with ``aligned_modbam``: the FASTA the BAM was aligned to, and
    # the caller's job directory, which alone holds modkit's work files (they
    # carry read names; never the system temp directory).
    reference_fasta: Path | None = None
    job_directory: Path | None = None
    modkit_filter_threshold: float = DEFAULT_MODKIT_FILTER_THRESHOLD
    min_mapq: int = DEFAULT_MIN_MAPQ


@dataclass(frozen=True, slots=True)
class _LoyferResources:
    markers: tuple[GenomicMarker, ...]
    atlas: AtlasUMatrix
    labels: Mapping[str, str]
    registered_marker_count: int
    excluded_incomplete_count: int
    collapsed_duplicate_count: int


@dataclass(frozen=True, slots=True)
class _LoyferHealthyRow:
    cell_type_id: str
    label: str
    donor_fractions: tuple[float, ...]
    reported_mean: float


@dataclass(frozen=True, slots=True)
class _LoyferHealthyTable:
    rows: tuple[_LoyferHealthyRow, ...]
    donor_ids: tuple[str, ...]
    source_id: str


_IDENTIFIER = TypeAdapter(Identifier)
# What a published result may not contain (``_model_safe``).  Labels and IDs
# read from Loyfer files reach the result, so the readers refuse them early.
_FORBIDDEN_OUTPUT_KEYS = frozenset({"fragment_digest", "path", "read_id", "sequence"})
_FORBIDDEN_OUTPUT_SUBSTRING = ".bam"


def _output_safe(value: str) -> bool:
    return value not in _FORBIDDEN_OUTPUT_KEYS and _FORBIDDEN_OUTPUT_SUBSTRING not in value


def _identifier(value: str, *, kind: str) -> str:
    normalized = re.sub(r"[^A-Za-z0-9_.:-]+", "-", value.strip())
    normalized = re.sub(r"-{2,}", "-", normalized).strip("-")
    try:
        # The models' own identifier rule (alphanumeric first character), so
        # a label that normalizes but cannot be a model ID is refused here.
        identifier = _IDENTIFIER.validate_python(normalized)
    except ValidationError:
        raise CellOriginPipelineError(f"{kind} cannot be normalized safely") from None
    if not (_output_safe(identifier) and _output_safe(value)):
        raise CellOriginPipelineError(f"{kind} cannot appear in a published result")
    return identifier


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    try:
        with path.open("rb") as handle:
            while chunk := handle.read(1024 * 1024):
                digest.update(chunk)
    except OSError as exc:
        raise CellOriginPipelineError("unable to checksum an input artifact") from exc
    return digest.hexdigest()


def _artifact(path: Path, artifact_id: str) -> DigestArtifact:
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise CellOriginPipelineError("unable to inspect an input artifact") from exc
    return DigestArtifact(
        artifact_id=artifact_id,
        sha256=_sha256(path),
        size_bytes=size,
    )


def _version_token(raw: str) -> str:
    match = re.search(r"[0-9]+(?:\.[0-9A-Za-z]+)+(?:[-+][0-9A-Za-z.-]+)?", raw)
    if match:
        return match.group(0)[:64]
    cleaned = re.sub(r"[^A-Za-z0-9_.+:-]+", "-", raw.strip()).strip("-")
    return (cleaned or "unknown")[:64]


def _tool_version(executable: str) -> tuple[bool, str | None]:
    try:
        completed = subprocess.run(
            [executable, "--version"],
            check=False,
            capture_output=True,
            text=True,
            errors="replace",  # Debian's samtools prints Latin-1 bytes
            timeout=10,
        )
    except (OSError, subprocess.SubprocessError):
        return False, None
    output = (completed.stdout or completed.stderr).strip().splitlines()
    return completed.returncode == 0 and bool(output), (
        _version_token(output[0]) if output else None
    )


def preflight(
    config: PipelineConfig,
    *,
    require_healthy: bool = False,
    require_modkit: bool = True,
    executable_finder: Callable[[str], str | None] = shutil.which,
    version_probe: Callable[[str], tuple[bool, str | None]] = _tool_version,
    modkit_resolver: Callable[[], PinnedTool] = resolve_modkit,
) -> PreflightReport:
    """Check tools and required local artifacts without exposing their paths.

    modkit is resolved only as the pinned, digest-checked toolchain, never
    from ``PATH``.
    """

    items: list[PreflightItem] = []
    for tool in ("minimap2", "samtools"):
        executable = executable_finder(tool)
        if executable is None:
            detail = f"{tool} is not installed or is not on PATH"
            items.append(
                PreflightItem(
                    item_id=f"software.{tool}",
                    ready=False,
                    blocking=True,
                    detail=detail,
                )
            )
            continue
        ready, version = version_probe(executable)
        items.append(
            PreflightItem(
                item_id=f"software.{tool}",
                ready=ready,
                blocking=True,
                version=version,
                detail=(
                    f"{tool} version detected"
                    if ready
                    else f"{tool} exists but its version could not be verified"
                ),
            )
        )

    try:
        modkit = modkit_resolver()
    except ToolProblem as problem:
        items.append(
            PreflightItem(
                item_id="software.modkit",
                ready=False,
                blocking=require_modkit,
                detail=f"{problem.code} ({problem.reason}): {problem.fix}",
            )
        )
    else:
        items.append(
            PreflightItem(
                item_id="software.modkit",
                ready=True,
                blocking=require_modkit,
                version=modkit.identity.version,
                detail="pinned modkit verified by version and binary digest",
            )
        )

    artifact_checks: list[tuple[str, Path | None]] = [
        ("atlas.marker-bed", config.marker_bed),
        ("atlas.marker-metadata", config.marker_metadata),
        ("atlas.u-matrix", config.atlas_u_matrix),
    ]
    if config.aligned_modbam is not None:
        artifact_checks.insert(0, ("reference.fasta", config.reference_fasta))
    for item_id, path in artifact_checks:
        ready = path is not None and path.is_file() and path.stat().st_size > 0
        items.append(
            PreflightItem(
                item_id=item_id,
                ready=ready,
                blocking=True,
                detail=(
                    f"{item_id} is available"
                    if ready
                    else f"{item_id} is missing or empty"
                ),
            )
        )

    input_count = sum(
        item is not None for item in (config.extract_tsv, config.aligned_modbam)
    )
    input_ready = input_count == 1 and (
        (config.extract_tsv is not None and config.extract_tsv.is_file())
        or (
            config.aligned_modbam is not None
            and config.aligned_modbam.is_file()
        )
    )
    items.append(
        PreflightItem(
            item_id="input.methylation-calls",
            ready=input_ready,
            blocking=True,
            detail=(
                "exactly one aligned modBAM or validated modkit extract is available"
                if input_ready
                else "provide exactly one existing aligned modBAM or validated modkit extract TSV"
            ),
        )
    )

    healthy_ready = (
        config.healthy_table is not None and config.healthy_table.is_file()
    )
    items.append(
        PreflightItem(
            item_id="reference.healthy-table-s8",
            ready=healthy_ready,
            blocking=require_healthy,
            detail=(
                "healthy Table S8 is available as contextual reference data"
                if healthy_ready
                else "provide a normalized healthy Table S8 file for the range chart"
            ),
        )
    )
    return PreflightReport(items=tuple(items))


def build_alignment_command_plan(
    source_modbam: Path,
    reference_index: Path,
    work_directory: Path,
    *,
    threads: int = 4,
) -> AlignmentCommandPlan:
    """Return, but never execute, a minimap2 alignment plan.

    ``samtools fastq -T`` serializes MM/ML/MN as FASTQ comments and minimap2
    ``-y`` copies those comments back into SAM auxiliary fields. The plan still
    requires post-run tag and sequence-orientation validation before modkit use.
    """

    if isinstance(threads, bool) or not isinstance(threads, int) or threads < 1:
        raise CellOriginPipelineError("threads must be a positive integer")
    fastq = work_directory / "reads.with-mod-tags.fastq"
    sam = work_directory / "aligned.with-mod-tags.sam"
    unsorted_bam = work_directory / "aligned.unsorted.bam"
    aligned_bam = work_directory / "aligned.sorted.bam"
    return AlignmentCommandPlan(
        steps=(
            CommandStep(
                name="export-fastq-with-tags",
                argv=(
                    "samtools",
                    "fastq",
                    "-T",
                    "MM,ML,MN",
                    "-n",
                    "-0",
                    str(fastq),
                    "-1",
                    os.devnull,
                    "-2",
                    os.devnull,
                    "-s",
                    os.devnull,
                    str(source_modbam),
                ),
                purpose="Copy primary Nanopore reads and modification tags to FASTQ comments.",
            ),
            CommandStep(
                name="align-hg38",
                argv=(
                    "minimap2",
                    "-a",
                    "-x",
                    "map-ont",
                    "-y",
                    "-t",
                    str(threads),
                    "-o",
                    str(sam),
                    str(reference_index),
                    str(fastq),
                ),
                purpose="Align to hg38 while restoring FASTQ comment tags to SAM.",
            ),
            CommandStep(
                name="sam-to-bam",
                argv=(
                    "samtools",
                    "view",
                    "-b",
                    "-o",
                    str(unsorted_bam),
                    str(sam),
                ),
                purpose="Convert SAM to BAM without a shell pipeline.",
            ),
            CommandStep(
                name="coordinate-sort",
                argv=(
                    "samtools",
                    "sort",
                    "-@",
                    str(threads),
                    "-o",
                    str(aligned_bam),
                    str(unsorted_bam),
                ),
                purpose="Coordinate-sort the aligned modBAM.",
            ),
            CommandStep(
                name="index",
                argv=("samtools", "index", str(aligned_bam)),
                purpose="Index the coordinate-sorted aligned modBAM.",
            ),
        ),
        postconditions=(
            "Compare input and output primary-read counts.",
            "Sample both orientations and verify MM, ML, and MN remain valid.",
            "Run modkit validate on the aligned modBAM before extraction.",
            "Do not delete the source modBAM or intermediates until checks pass.",
        ),
    )


def _strict_dict_reader(
    path: Path,
    *,
    expected_header: Sequence[str],
    max_rows: int,
) -> tuple[dict[str, str], ...]:
    try:
        with path.open("r", encoding="utf-8", newline="") as handle:
            reader = csv.DictReader(handle, delimiter="\t")
            if tuple(reader.fieldnames or ()) != tuple(expected_header):
                raise CellOriginPipelineError(
                    "a Loyfer resource header does not match its registered schema"
                )
            rows: list[dict[str, str]] = []
            for row_number, row in enumerate(reader, start=1):
                if row_number > max_rows:
                    raise CellOriginPipelineError(
                        f"a Loyfer resource exceeds the row cap of {max_rows}"
                    )
                if None in row or any(value is None for value in row.values()):
                    raise CellOriginPipelineError(
                        "a Loyfer resource contains a malformed row"
                    )
                rows.append(dict(row))
    except (OSError, UnicodeError, csv.Error) as exc:
        raise CellOriginPipelineError("unable to parse a Loyfer resource") from exc
    if not rows:
        raise CellOriginPipelineError("a Loyfer resource contains no records")
    return tuple(rows)


# Identity used when a shared reader runs the downstream loader on its own
# file (registration has no atlas or source ID yet; neither changes validity).
_SELF_CHECK_ID = "self-check"


def _derived_atlas_matrix(
    cell_ids: tuple[str, ...],
    complete: Sequence[tuple[str, tuple[str, ...]]],
    *,
    atlas_id: str,
    source_ids: tuple[str, ...],
) -> AtlasUMatrix:
    """Load the complete atlas rows through ``load_loyfer_atlas_u_matrix``."""

    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, delimiter="\t", lineterminator="\n")
    writer.writerow(("marker_id", *cell_ids))
    writer.writerows((region, *values) for region, values in complete)
    buffer.seek(0)
    try:
        return load_loyfer_atlas_u_matrix(
            buffer,
            columns=AtlasUColumns(
                marker_id="marker_id",
                cell_type_columns=tuple((item, item) for item in cell_ids),
            ),
            atlas_id=atlas_id,
            source_ids=source_ids,
            expected_marker_ids=tuple(region for region, _ in complete),
            expected_cell_type_ids=cell_ids,
            max_rows=LOYFER_MAX_ROWS,
        )
    except CellOriginInputError as exc:
        raise CellOriginPipelineError(str(exc)) from exc


def _derived_markers(
    metadata_groups: Mapping[str, Sequence[Mapping[str, str]]],
    regions: Sequence[str],
    *,
    atlas_id: str,
    source_ids: tuple[str, ...],
    expected_cell_type_ids: Sequence[str] | None,
) -> tuple[GenomicMarker, ...]:
    """Load the marker rows of ``regions`` through ``load_marker_bed``."""

    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, delimiter="\t", lineterminator="\n")
    for region in regions:
        rows = metadata_groups[region]
        row = rows[0]
        targets = {item["target"] for item in rows}
        writer.writerow(
            (
                row["#chr"],
                row["start"],
                row["end"],
                region,
                (
                    _identifier(next(iter(targets)), kind="marker target")
                    if len(targets) == 1
                    else "multi-target"
                ),
            )
        )
    buffer.seek(0)
    try:
        return load_marker_bed(
            buffer,
            columns=MarkerBedColumns(
                chromosome=0,
                start0=1,
                end0=2,
                marker_id=3,
                target_cell_type_id=4,
            ),
            atlas_id=atlas_id,
            source_ids=source_ids,
            expected_marker_ids=tuple(regions),
            expected_cell_type_ids=expected_cell_type_ids,
            max_rows=LOYFER_MAX_ROWS,
        )
    except CellOriginInputError as exc:
        raise CellOriginPipelineError(str(exc)) from exc


def read_marker_regions(
    path: Path, *, max_rows: int = LOYFER_MAX_ROWS
) -> tuple[str, ...]:
    """Parse ``Regions.U250`` (headerless chrom/start/end BED) into region IDs.

    The one parser of this file: the pipeline and method-asset registration
    (``traceback_runner.references``) both call it.
    """

    regions: list[str] = []
    try:
        with path.open("r", encoding="utf-8", newline="") as handle:
            # QUOTE_NONE: fields are read exactly as the prefilter's
            # read_bed_regions reads them (a quote is part of the value).
            reader = csv.reader(handle, delimiter="\t", quoting=csv.QUOTE_NONE)
            for row_number, row in enumerate(reader, start=1):
                if row_number > max_rows:
                    raise CellOriginPipelineError(
                        f"marker BED exceeds the row cap of {max_rows}"
                    )
                if len(row) != 3:
                    raise CellOriginPipelineError(
                        "marker BED must contain exactly three columns"
                    )
                chromosome, start, end = row
                try:
                    start0, end0 = int(start), int(end)
                except ValueError as exc:
                    raise CellOriginPipelineError(
                        "marker BED contains a non-integer coordinate"
                    ) from exc
                if start != str(start0) or end != str(end0) or end0 <= start0:
                    raise CellOriginPipelineError(
                        "marker BED contains an invalid half-open interval"
                    )
                region = f"{chromosome}:{start0}-{end0}"
                try:
                    # The GenomicMarker rules (chromosome name, start0 >= 0)
                    # hold for every row, not only for the usable ones.
                    GenomicMarker(
                        marker_id=region,
                        chromosome=chromosome,
                        start0=start0,
                        end0=end0,
                        target_cell_type_id="registration",
                        atlas_id="registration",
                        source_ids=("registration",),
                    )
                except ValidationError as exc:
                    raise CellOriginPipelineError(
                        "marker BED contains an unsupported chromosome or a negative start"
                    ) from exc
                regions.append(region)
    except (OSError, UnicodeError, csv.Error) as exc:
        raise CellOriginPipelineError("unable to parse marker BED") from exc
    if len(regions) != len(set(regions)):
        raise CellOriginPipelineError("marker BED contains duplicate regions")
    try:
        # The extraction step reads the same file with its own parser.
        read_bed_regions(path)
    except (OSError, UnicodeError, ValueError) as exc:
        raise CellOriginPipelineError("the marker region BED could not be read") from exc
    return tuple(regions)


def read_marker_metadata(
    path: Path, *, max_rows: int = LOYFER_MAX_ROWS
) -> tuple[tuple[dict[str, str], ...], dict[str, list[dict[str, str]]]]:
    """Parse ``Markers.U250`` into its rows and rows grouped by region ID.

    Rows sharing a region ID must agree on coordinates and direction.  The one
    parser of this file (pipeline and method-asset registration).
    """

    metadata_rows = _strict_dict_reader(
        path,
        expected_header=MARKER_METADATA_HEADER,
        max_rows=max_rows,
    )
    metadata_groups: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in metadata_rows:
        metadata_groups[row["region"]].append(row)
    for rows in metadata_groups.values():
        first = rows[0]
        if any(
            (
                row["#chr"],
                row["start"],
                row["end"],
                row["direction"],
            )
            != (
                first["#chr"],
                first["start"],
                first["end"],
                first["direction"],
            )
            for row in rows[1:]
        ):
            raise CellOriginPipelineError(
                "duplicate Loyfer region IDs disagree on coordinates or direction"
            )
    # The downstream marker loader on this file's own rows (the analysis runs
    # it again on the usable subset, with the atlas's cell types).
    _derived_markers(
        metadata_groups,
        tuple(metadata_groups),
        atlas_id=_SELF_CHECK_ID,
        source_ids=(_SELF_CHECK_ID,),
        expected_cell_type_ids=None,
    )
    return metadata_rows, dict(metadata_groups)


@dataclass(frozen=True, slots=True)
class AtlasU250Table:
    """``Atlas.U250`` checked on its own; cross-file checks come later.

    ``rows`` holds every data row in file order (duplicates included),
    ``by_region`` the first row of each region, and ``complete`` the
    ``(region, raw values)`` of each region with a value for every cell type.
    """

    raw_cell_labels: tuple[str, ...]
    cell_ids: tuple[str, ...]
    rows: tuple[dict[str, str], ...]
    by_region: Mapping[str, dict[str, str]]
    order: tuple[str, ...]
    complete: tuple[tuple[str, tuple[str, ...]], ...]
    incomplete: int


def read_atlas_u250(path: Path, *, max_rows: int = LOYFER_MAX_ROWS) -> AtlasU250Table:
    """Parse ``Atlas.U250`` and check it on its own.

    The header, cell-type label normalization and collisions, the row cap,
    malformed rows, the ``U`` direction, duplicate-region agreement and the
    U fractions of complete rows.  The one parser of this file (pipeline and
    method-asset registration).
    """

    try:
        with path.open("r", encoding="utf-8", newline="") as handle:
            reader = csv.DictReader(handle, delimiter="\t")
            header = tuple(reader.fieldnames or ())
            if header[:8] != ATLAS_METADATA_HEADER or len(header) <= 8:
                raise CellOriginPipelineError(
                    "Atlas.U250 header does not match the registered schema"
                )
            raw_cell_labels = header[8:]
            cell_ids = tuple(
                _identifier(value, kind="cell type") for value in raw_cell_labels
            )
            # "marker_id" keys the derived atlas table, so no cell type may use it.
            if len(cell_ids) != len(set(cell_ids)) or "marker_id" in cell_ids:
                raise CellOriginPipelineError(
                    "cell type labels collide after identifier normalization"
                )
            rows: list[dict[str, str]] = []
            atlas_by_region: dict[str, dict[str, str]] = {}
            atlas_order: list[str] = []
            for row_number, row in enumerate(reader, start=1):
                if row_number > max_rows:
                    raise CellOriginPipelineError(
                        f"Atlas.U250 exceeds the row cap of {max_rows}"
                    )
                if None in row or any(value is None for value in row.values()):
                    raise CellOriginPipelineError(
                        "Atlas.U250 contains a malformed row"
                    )
                if row["direction"] != "U":
                    raise CellOriginPipelineError(
                        "Loyfer marker resources disagree on coordinates, target, or direction"
                    )
                rows.append(row)
                region = row["name"]
                previous = atlas_by_region.get(region)
                if previous is not None:
                    comparable_fields = (
                        "chr",
                        "start",
                        "end",
                        "startCpG",
                        "endCpG",
                        "name",
                        "direction",
                        *raw_cell_labels,
                    )
                    if any(previous[key] != row[key] for key in comparable_fields):
                        raise CellOriginPipelineError(
                            "duplicate Atlas.U250 regions disagree on U values"
                        )
                    continue
                atlas_by_region[region] = dict(row)
                atlas_order.append(region)
    except (OSError, UnicodeError, csv.Error) as exc:
        raise CellOriginPipelineError("unable to parse Atlas.U250") from exc
    complete: list[tuple[str, tuple[str, ...]]] = []
    incomplete = 0
    for region in atlas_order:
        row = atlas_by_region[region]
        values = tuple(row[label] for label in raw_cell_labels)
        if any(value in {"", "NA"} for value in values):
            incomplete += 1
            continue
        try:
            parsed = tuple(float(value) for value in values)
        except ValueError as exc:
            raise CellOriginPipelineError(
                "Atlas.U250 contains a nonnumeric U fraction"
            ) from exc
        if any(not np.isfinite(value) or not 0.0 <= value <= 1.0 for value in parsed):
            raise CellOriginPipelineError("Atlas.U250 contains an invalid U fraction")
        complete.append((region, values))
    if not complete:
        raise CellOriginPipelineError(
            "Atlas.U250 has no complete marker rows across all cell types"
        )
    # The downstream atlas loader on exactly the rows the analysis loads.
    _derived_atlas_matrix(
        cell_ids, complete, atlas_id=_SELF_CHECK_ID, source_ids=(_SELF_CHECK_ID,)
    )
    return AtlasU250Table(
        raw_cell_labels=raw_cell_labels,
        cell_ids=cell_ids,
        rows=tuple(rows),
        by_region=atlas_by_region,
        order=tuple(atlas_order),
        complete=tuple(complete),
        incomplete=incomplete,
    )


def _load_loyfer_resources(config: PipelineConfig) -> _LoyferResources:
    metadata_rows, metadata_groups = read_marker_metadata(config.marker_metadata)
    metadata_by_region = {
        region: rows[0] for region, rows in metadata_groups.items()
    }
    collapsed_duplicates = len(metadata_rows) - len(metadata_by_region)
    regions = read_marker_regions(config.marker_bed)
    if set(regions) != set(metadata_by_region):
        raise CellOriginPipelineError(
            "marker BED and Markers.U250 region IDs do not match exactly"
        )

    table = read_atlas_u250(config.atlas_u_matrix)
    raw_cell_labels = table.raw_cell_labels
    cell_ids = table.cell_ids
    labels = dict(zip(cell_ids, raw_cell_labels, strict=True))
    atlas_by_region = table.by_region
    incomplete = table.incomplete
    for row in table.rows:
        region = row["name"]
        metadata_rows_for_region = metadata_groups.get(region)
        if metadata_rows_for_region is None:
            raise CellOriginPipelineError(
                "Atlas.U250 contains an unregistered marker"
            )
        metadata = metadata_rows_for_region[0]
        registered_targets = {
            item["target"] for item in metadata_rows_for_region
        }
        if (
            row["chr"] != metadata["#chr"]
            or row["start"] != metadata["start"]
            or row["end"] != metadata["end"]
            or row["target"] not in registered_targets
            or row["direction"] != metadata["direction"]
        ):
            raise CellOriginPipelineError(
                "Loyfer marker resources disagree on coordinates, target, or direction"
            )
    if set(atlas_by_region) != set(regions):
        raise CellOriginPipelineError(
            "Atlas.U250 and marker BED region IDs do not match exactly"
        )
    usable_regions = tuple(region for region, _ in table.complete)
    atlas = _derived_atlas_matrix(
        cell_ids,
        table.complete,
        atlas_id=config.atlas_id,
        source_ids=(config.atlas_source_id,),
    )
    markers = _derived_markers(
        metadata_groups,
        usable_regions,
        atlas_id=config.atlas_id,
        source_ids=(config.atlas_source_id,),
        expected_cell_type_ids=(*cell_ids, "multi-target"),
    )
    return _LoyferResources(
        markers=markers,
        atlas=atlas,
        labels=labels,
        registered_marker_count=len(metadata_rows),
        excluded_incomplete_count=incomplete,
        collapsed_duplicate_count=collapsed_duplicates,
    )


def _interval_index(
    markers: Sequence[GenomicMarker],
) -> dict[str, tuple[tuple[int, ...], tuple[int, ...], tuple[GenomicMarker, ...]]]:
    by_chromosome: dict[str, list[GenomicMarker]] = defaultdict(list)
    for marker in markers:
        by_chromosome[marker.chromosome].append(marker)
    index = {}
    for chromosome, values in by_chromosome.items():
        ordered = tuple(sorted(values, key=lambda item: (item.start0, item.end0)))
        starts = tuple(item.start0 for item in ordered)
        prefix_max: list[int] = []
        maximum = -1
        for item in ordered:
            maximum = max(maximum, item.end0)
            prefix_max.append(maximum)
        index[chromosome] = (starts, tuple(prefix_max), ordered)
    return index


def _matching_marker(
    call: Any,
    index: Mapping[
        str, tuple[tuple[int, ...], tuple[int, ...], tuple[GenomicMarker, ...]]
    ],
) -> tuple[GenomicMarker | None, bool]:
    chromosome_index = index.get(call.chromosome)
    if chromosome_index is None:
        return None, False
    starts, prefix_max, markers = chromosome_index
    cursor = bisect_right(starts, call.position0) - 1
    matches: list[GenomicMarker] = []
    while cursor >= 0 and prefix_max[cursor] > call.position0:
        marker = markers[cursor]
        if marker.start0 <= call.position0 < marker.end0:
            matches.append(marker)
            if len(matches) > 1:
                return None, True
        cursor -= 1
    return (matches[0], False) if matches else (None, False)


def _partitioned_classification(
    calls: Sequence[Any],
    markers: Sequence[GenomicMarker],
    *,
    maximum_groups: int,
    maximum_cpgs_per_group: int,
) -> UxmClassificationResult:
    """Use an interval index, then delegate exact UXM semantics marker by marker."""

    index = _interval_index(markers)
    by_marker: dict[str, list[Any]] = defaultdict(list)
    marker_by_id = {marker.marker_id: marker for marker in markers}
    outside = 0
    ambiguous = 0
    invalid_fragments: set[str] = set()
    assigned: list[tuple[str, Any]] = []
    for call in calls:
        marker, is_ambiguous = _matching_marker(call, index)
        if is_ambiguous:
            ambiguous += 1
            invalid_fragments.add(call.fragment_digest)
        elif marker is None:
            outside += 1
        else:
            assigned.append((marker.marker_id, call))
    for marker_id, call in assigned:
        if call.fragment_digest not in invalid_fragments:
            by_marker[marker_id].append(call)

    observations = []
    diagnostic_values: dict[str, int] = defaultdict(int)
    remaining_groups = maximum_groups
    stop_reason = UxmStopReason.COMPLETE_INPUT
    marker_order = {marker.marker_id: index for index, marker in enumerate(markers)}
    for marker_id in sorted(by_marker, key=marker_order.__getitem__):
        if remaining_groups < 1:
            stop_reason = UxmStopReason.GROUP_CAP
            break
        marker_calls = by_marker[marker_id]
        partial = classify_uxm_calls(
            marker_calls,
            (marker_by_id[marker_id],),
            maximum_calls=len(marker_calls) + 1,
            maximum_groups=remaining_groups,
            maximum_cpgs_per_group=maximum_cpgs_per_group,
        )
        observations.extend(partial.observations)
        values = asdict(partial.diagnostics)
        for key, value in values.items():
            if key not in {"stop_reason", "partial_input"}:
                diagnostic_values[key] += int(value)
        remaining_groups -= partial.diagnostics.fragment_marker_group_count
        if partial.diagnostics.stop_reason == UxmStopReason.GROUP_CAP:
            stop_reason = UxmStopReason.GROUP_CAP
            break

    marker_counts, marker_weights = aggregate_marker_counts(
        observations, marker_order=markers
    )
    diagnostics = UxmDiagnostics(
        inspected_calls=len(calls),
        retained_unique_cpg_calls=diagnostic_values[
            "retained_unique_cpg_calls"
        ],
        duplicate_cpg_calls=diagnostic_values["duplicate_cpg_calls"],
        excluded_failed_calls=diagnostic_values["excluded_failed_calls"],
        excluded_non_c_calls=diagnostic_values["excluded_non_c_calls"],
        excluded_unsupported_modification_calls=diagnostic_values[
            "excluded_unsupported_modification_calls"
        ],
        excluded_invalid_calls=diagnostic_values["excluded_invalid_calls"],
        excluded_outside_marker_calls=outside,
        excluded_ambiguous_marker_calls=ambiguous,
        excluded_unknown_marker_calls=diagnostic_values[
            "excluded_unknown_marker_calls"
        ],
        excluded_conflicting_duplicate_groups=diagnostic_values[
            "excluded_conflicting_duplicate_groups"
        ],
        excluded_oversized_groups=diagnostic_values[
            "excluded_oversized_groups"
        ],
        fragment_marker_group_count=diagnostic_values[
            "fragment_marker_group_count"
        ],
        classified_fragment_marker_count=len(observations),
        excluded_fewer_than_four_cpgs=diagnostic_values[
            "excluded_fewer_than_four_cpgs"
        ],
        stop_reason=stop_reason,
        partial_input=stop_reason != UxmStopReason.COMPLETE_INPUT,
    )
    return UxmClassificationResult(
        observations=tuple(observations),
        marker_counts=marker_counts,
        marker_weights=marker_weights,
        diagnostics=diagnostics,
    )


def _filtered_atlas(
    atlas: AtlasUMatrix, marker_ids: Sequence[str]
) -> AtlasUMatrix:
    selected = set(marker_ids)
    rows = tuple(row for row in atlas.rows if row.marker_id in selected)
    if len(rows) != len(selected):
        raise CellOriginPipelineError(
            "classified marker IDs do not match the registered atlas"
        )
    return AtlasUMatrix(
        atlas_id=atlas.atlas_id,
        method=atlas.method,
        cell_type_ids=atlas.cell_type_ids,
        rows=rows,
        source_ids=atlas.source_ids,
    )


def _load_healthy(
    config: PipelineConfig,
    cell_type_ids: Sequence[str],
    labels: Mapping[str, str],
) -> _LoyferHealthyTable | None:
    if config.healthy_table is None:
        return None
    raw_to_id = {label: cell_id for cell_id, label in labels.items()}
    rows: list[_LoyferHealthyRow] = []
    donor_ids: tuple[str, ...]
    suffix = config.healthy_table.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise CellOriginPipelineError(
                "openpyxl is required to read Loyfer Supplementary Table S8"
            ) from exc
        try:
            workbook = load_workbook(
                config.healthy_table, read_only=True, data_only=True
            )
            sheet_name = config.healthy_sheet or "Table S8"
            worksheet = workbook[sheet_name]
            title = worksheet.cell(1, 1).value
            if (
                not isinstance(title, str)
                or "Supplementary Table S8" not in title
                or "Fragment-level deconvolution" not in title
            ):
                raise CellOriginPipelineError(
                    "workbook sheet is not registered Loyfer fragment-level Table S8"
                )
            if (
                worksheet.cell(3, 1).value != "tissue"
                or worksheet.cell(3, 2).value != "Plasma (n=23)"
            ):
                raise CellOriginPipelineError(
                    "Loyfer Table S8 plasma header does not match the registered schema"
                )
            raw_donor_ids = tuple(
                worksheet.cell(3, column).value for column in range(3, 26)
            )
            if len(raw_donor_ids) != 23 or any(
                not isinstance(item, str) or not item
                for item in raw_donor_ids
            ):
                raise CellOriginPipelineError(
                    "Loyfer Table S8 must contain exactly 23 plasma donor columns C:Y"
                )
            donor_ids = tuple(
                _identifier(item, kind="plasma donor")
                for item in raw_donor_ids
                if isinstance(item, str)
            )
            for row_number in range(4, worksheet.max_row + 1):
                raw_label = worksheet.cell(row_number, 1).value
                if raw_label is None:
                    if rows:
                        break
                    continue
                if not isinstance(raw_label, str) or not raw_label:
                    raise CellOriginPipelineError(
                        "Loyfer Table S8 contains an invalid tissue label"
                    )
                raw_mean = worksheet.cell(row_number, 2).value
                raw_values = tuple(
                    worksheet.cell(row_number, column).value
                    for column in range(3, 26)
                )
                if (
                    isinstance(raw_mean, bool)
                    or not isinstance(raw_mean, (int, float))
                    or any(
                        isinstance(value, bool)
                        or not isinstance(value, (int, float))
                        for value in raw_values
                    )
                ):
                    raise CellOriginPipelineError(
                        "Loyfer Table S8 plasma values must be numeric"
                    )
                mean = float(raw_mean)
                values = tuple(float(value) for value in raw_values)
                if any(
                    not np.isfinite(value) or not 0.0 <= value <= 1.0
                    for value in (mean, *values)
                ):
                    raise CellOriginPipelineError(
                        "Loyfer Table S8 plasma fractions must be within [0, 1]"
                    )
                if not np.isclose(
                    mean, float(np.mean(values)), rtol=0.0, atol=5e-6
                ):
                    raise CellOriginPipelineError(
                        "Loyfer Table S8 reported mean does not match its 23 donors"
                    )
                cell_id = (
                    "Other"
                    if raw_label == "Other"
                    else raw_to_id.get(
                        raw_label, _identifier(raw_label, kind="cell type")
                    )
                )
                rows.append(
                    _LoyferHealthyRow(
                        cell_type_id=cell_id,
                        label=raw_label,
                        donor_fractions=values,
                        reported_mean=mean,
                    )
                )
            workbook.close()
        except (OSError, KeyError, ValueError) as exc:
            if isinstance(exc, CellOriginPipelineError):
                raise
            raise CellOriginPipelineError(
                "unable to read Loyfer Supplementary Table S8"
            ) from exc
    else:
        delimiter = "\t" if suffix in {".tsv", ".txt"} else ","
        try:
            with config.healthy_table.open(
                "r", encoding="utf-8", newline=""
            ) as handle:
                reader = csv.DictReader(handle, delimiter=delimiter)
                header = tuple(reader.fieldnames or ())
                if (
                    len(header) < 4
                    or header[0] != "cell_type_id"
                    or header[1] != "mean"
                ):
                    raise CellOriginPipelineError(
                        "normalized Loyfer Table S8 must start with cell_type_id, mean, and donor columns"
                    )
                donor_columns = header[2:]
                donor_ids = tuple(
                    _identifier(item, kind="plasma donor")
                    for item in donor_columns
                )
                divisor = (
                    100.0
                    if config.healthy_fraction_unit == FractionUnit.PERCENT
                    else 1.0
                )
                for record in reader:
                    raw_label = record["cell_type_id"]
                    values = tuple(
                        float(record[column]) / divisor
                        for column in donor_columns
                    )
                    mean = float(record["mean"]) / divisor
                    cell_id = (
                        "Other"
                        if raw_label == "Other"
                        else raw_to_id.get(
                            raw_label, _identifier(raw_label, kind="cell type")
                        )
                    )
                    rows.append(
                        _LoyferHealthyRow(
                            cell_type_id=cell_id,
                            label=raw_label,
                            donor_fractions=values,
                            reported_mean=mean,
                        )
                    )
        except (OSError, UnicodeError, csv.Error, KeyError, ValueError) as exc:
            if isinstance(exc, CellOriginPipelineError):
                raise
            raise CellOriginPipelineError(
                "unable to read normalized Loyfer Table S8"
            ) from exc
    if not rows:
        raise CellOriginPipelineError("Loyfer Table S8 contains no plasma rows")
    row_ids = tuple(row.cell_type_id for row in rows)
    if len(row_ids) != len(set(row_ids)):
        raise CellOriginPipelineError(
            "Loyfer Table S8 contains duplicate tissue categories"
        )
    unknown = set(row_ids) - set(cell_type_ids) - {"Other"}
    if unknown:
        raise CellOriginPipelineError(
            "Loyfer Table S8 contains tissue IDs absent from the atlas"
        )
    if "Other" not in row_ids:
        raise CellOriginPipelineError(
            "Loyfer Table S8 must include its published Other category"
        )
    return _LoyferHealthyTable(
        rows=tuple(rows),
        donor_ids=donor_ids,
        source_id=config.healthy_source_id,
    )


def _healthy_chart_rows(
    healthy: _LoyferHealthyTable,
    estimates: Mapping[str, float],
    intervals: Mapping[str, tuple[float | None, float | None]],
    labels: Mapping[str, str],
) -> tuple[HealthyRangeChartRow, ...]:
    direct_ids = {
        row.cell_type_id for row in healthy.rows if row.cell_type_id != "Other"
    }
    rows = []
    for row in healthy.rows:
        values = np.asarray(row.donor_fractions, dtype=np.float64)
        q1, median, q3 = np.quantile(
            values, (0.25, 0.5, 0.75), method="linear"
        )
        if row.cell_type_id == "Other":
            estimate = sum(
                fraction
                for cell_id, fraction in estimates.items()
                if cell_id not in direct_ids
            )
            lower = upper = None
        else:
            estimate = estimates[row.cell_type_id]
            lower, upper = intervals[row.cell_type_id]
        minimum = float(np.min(values))
        maximum = float(np.max(values))
        classification: Literal["below", "within", "above"] = (
            "below"
            if estimate < minimum
            else "above"
            if estimate > maximum
            else "within"
        )
        rows.append(
            HealthyRangeChartRow(
                cell_type_id=row.cell_type_id,
                label=(
                    row.label
                    if row.cell_type_id == "Other"
                    else labels.get(row.cell_type_id, row.label)
                ),
                sample_fraction=estimate,
                sample_percent=estimate * 100.0,
                sample_lower_percent=(
                    lower * 100.0 if lower is not None else None
                ),
                sample_upper_percent=(
                    upper * 100.0 if upper is not None else None
                ),
                healthy_min_percent=minimum * 100.0,
                healthy_q1_percent=float(q1) * 100.0,
                healthy_median_percent=float(median) * 100.0,
                healthy_q3_percent=float(q3) * 100.0,
                healthy_max_percent=maximum * 100.0,
                healthy_mean_percent=row.reported_mean * 100.0,
                classification=classification,
                uncertainty_available=lower is not None,
            )
        )
    return tuple(
        sorted(rows, key=lambda item: (-item.sample_fraction, item.label))
    )


def _model_safe(result: CellOriginResult) -> bool:
    payload = result.model_dump(mode="json")
    serialized = json.dumps(payload, sort_keys=True)
    forbidden_keys = _FORBIDDEN_OUTPUT_KEYS

    def keys(value: Any) -> Iterable[str]:
        if isinstance(value, dict):
            for key, item in value.items():
                yield key
                yield from keys(item)
        elif isinstance(value, list):
            for item in value:
                yield from keys(item)

    return (
        not forbidden_keys.intersection(keys(payload))
        and _FORBIDDEN_OUTPUT_SUBSTRING not in serialized
    )


def _atomic_write(path: Path, bundle: CellOriginResultBundle) -> str:
    payload = bundle.model_dump_json(indent=2).encode("utf-8") + b"\n"
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="wb",
            dir=path.parent,
            prefix=f".{path.name}.",
            suffix=".tmp",
            delete=False,
        ) as handle:
            temporary = Path(handle.name)
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())
        CellOriginResultBundle.model_validate_json(temporary.read_bytes())
        os.replace(temporary, path)
    except (OSError, ValueError) as exc:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
        raise CellOriginPipelineError(
            "result bundle could not be validated and atomically published"
        ) from exc
    return hashlib.sha256(payload).hexdigest()


def _native_probability(raw: str, call_code: str) -> str:
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        parsed = None
    if isinstance(parsed, dict) and call_code in parsed:
        return str(parsed[call_code])
    pairs = dict(
        re.findall(
            r"([A-Za-z])\s*[:=]\s*([0-9]+(?:\.[0-9]+)?(?:[eE][+-]?[0-9]+)?)",
            raw,
        )
    )
    if call_code in pairs:
        return pairs[call_code]
    raise CellOriginPipelineError(
        "installed modkit emitted an unsupported base_probs encoding; normalize the extract explicitly"
    )


def _normalize_native_modkit(
    source: Path,
    destination: Path,
    *,
    maximum_calls: int = DEFAULT_MAXIMUM_CALLS,
) -> int:
    try:
        with source.open("r", encoding="utf-8", newline="") as input_handle:
            return _normalize_native_stream(
                input_handle, destination, maximum_calls=maximum_calls
            )
    except OSError as exc:
        raise CellOriginPipelineError(
            "unable to normalize modkit extract output"
        ) from exc


def _normalize_native_stream(
    input_handle: Iterable[str],
    destination: Path,
    *,
    maximum_calls: int,
) -> int:
    """Normalize native ``extract calls`` rows; refuse past ``maximum_calls``.

    The cap is checked row by row while streaming, so a caller reading modkit's
    stdout stops it as soon as the cap is passed.  Returns the row count.
    """

    if isinstance(maximum_calls, bool) or maximum_calls < 1:
        raise CellOriginPipelineError("the call cap must be a positive integer")
    rows = 0
    try:
        reader = csv.DictReader(input_handle, delimiter="\t")
        header = frozenset(reader.fieldnames or ())
        if not NATIVE_MODKIT_REQUIRED_COLUMNS.issubset(header):
            raise CellOriginPipelineError(
                "installed modkit emitted an unsupported extract schema"
            )
        probability_column = next(
            (
                name
                for name in ("modified_probability", "call_prob")
                if name in header
            ),
            None,
        )
        if probability_column is None and "base_probs" not in header:
            raise CellOriginPipelineError(
                "modkit extract lacks call_prob, modified_probability, or base_probs"
            )
        with destination.open("w", encoding="utf-8", newline="") as output_handle:
            writer = csv.DictWriter(
                output_handle,
                fieldnames=NORMALIZED_MODKIT_COLUMNS.declared(),
                delimiter="\t",
                lineterminator="\n",
            )
            writer.writeheader()
            for row_number, row in enumerate(reader, start=1):
                if row_number > maximum_calls:
                    raise CallCapExceeded(
                        "modkit extract exceeds the configured call cap"
                    )
                native_call_code = row["call_code"]
                call_code = "C" if native_call_code == "-" else native_call_code
                writer.writerow(
                    {
                        "read_id": row["read_id"],
                        "chrom": row["chrom"],
                        "ref_position": row["ref_position"],
                        "mod_strand": row["mod_strand"],
                        "modified_primary_base": row["modified_primary_base"],
                        "call_code": call_code,
                        "modified_probability": (
                            row[probability_column]
                            if probability_column is not None
                            else _native_probability(
                                row["base_probs"], native_call_code
                            )
                        ),
                        "fail": row["fail"].lower(),
                    }
                )
                rows = row_number
    except (OSError, UnicodeError, csv.Error, KeyError) as exc:
        raise CellOriginPipelineError(
            "unable to normalize modkit extract output"
        ) from exc
    return rows


def modkit_extract_arguments(
    *,
    reference_fasta: Path,
    include_bed: Path,
    filter_threshold: float,
    modbam: Path,
) -> tuple[str, ...]:
    """The locked ``modkit extract calls`` argv after the executable.

    Output goes to stdout (``-``) so the call cap is enforced while streaming;
    the filter threshold is always explicit, never modkit's estimate.
    """

    if not 0.0 < filter_threshold < 1.0:
        raise CellOriginPipelineError("the modkit filter threshold must be in (0, 1)")
    return (
        "extract",
        "calls",
        "--reference",
        str(reference_fasta),
        "--cpg",
        "--include-bed",
        str(include_bed),
        "--mapped-only",
        "--filter-threshold",
        repr(float(filter_threshold)),
        "--suppress-progress",
        str(modbam),
        "-",
    )


def _modkit_work_directory(config: PipelineConfig) -> Path:
    if config.job_directory is None:
        raise CellOriginPipelineError(
            "a job directory is required to extract calls from an aligned modBAM"
        )
    return config.job_directory.absolute() / MODKIT_WORK_DIRECTORY


def _remove_work_directory(work: Path) -> None:
    if work.is_symlink() or work.is_file():
        work.unlink()
    elif work.is_dir():
        shutil.rmtree(work)


def _extract_modbam(
    config: PipelineConfig,
    *,
    modkit: PinnedTool,
    timeout_seconds: float = MODKIT_TIMEOUT_SECONDS,
) -> tuple[Path, PrefilterCounts]:
    """Pre-filter, then stream ``modkit extract calls`` into a bounded TSV.

    Every file it writes is under ``<job directory>/modkit-work``; the caller
    removes that directory when the run ends.
    """

    if config.aligned_modbam is None:
        raise CellOriginPipelineError("aligned modBAM input is not configured")
    if config.reference_fasta is None:
        raise CellOriginPipelineError(
            "a reference FASTA is required to extract calls from an aligned modBAM"
        )
    # modkit runs with cwd=work, so every path it is given is absolute.
    reference_fasta = config.reference_fasta.absolute()
    marker_bed = config.marker_bed.absolute()
    work = _modkit_work_directory(config)
    _remove_work_directory(work)
    work.parent.mkdir(parents=True, exist_ok=True)
    # Private before anything holding read names is written into it.
    work.mkdir(mode=0o700)
    work.chmod(0o700)
    temporary = work / "tmp"
    temporary.mkdir(mode=0o700)
    try:
        regions = read_bed_regions(marker_bed)
    except (OSError, UnicodeError, ValueError) as exc:
        raise CellOriginPipelineError("the marker region BED could not be read") from exc
    filtered = work / "prefiltered.bam"
    try:
        counts = write_prefiltered_bam(
            config.aligned_modbam, filtered, regions, min_mapq=config.min_mapq
        )
    except (OSError, ValueError) as exc:
        raise CellOriginPipelineError(
            "the aligned modBAM could not be pre-filtered"
        ) from exc
    if counts.written == 0:
        raise NoMarkerAlignments(
            "no alignment passes the pre-filter and overlaps a marker region"
        )
    normalized = work / "modkit.normalized.tsv"
    arguments = modkit_extract_arguments(
        reference_fasta=reference_fasta,
        include_bed=marker_bed,
        filter_threshold=config.modkit_filter_threshold,
        modbam=filtered,
    )
    environment = {**os.environ, "TMPDIR": str(temporary)}
    with (work / "modkit.stderr.log").open("wb") as stderr_log:
        try:
            process = exec_pinned(
                modkit,
                arguments,
                stdin=subprocess.DEVNULL,
                stdout=subprocess.PIPE,
                stderr=stderr_log,
                cwd=work,
                env=environment,
            )
        except OSError as exc:
            raise CellOriginPipelineError("modkit extraction could not run") from exc
        timed_out = threading.Event()

        def _on_timeout() -> None:
            timed_out.set()
            kill_process_group(process)

        watchdog = threading.Timer(timeout_seconds, _on_timeout)
        watchdog.daemon = True
        watchdog.start()
        try:
            assert process.stdout is not None
            _normalize_native_stream(
                io.TextIOWrapper(process.stdout, encoding="utf-8", newline=""),
                normalized,
                maximum_calls=config.maximum_calls,
            )
            returncode = process.wait()
        except CellOriginPipelineError as exc:
            # The stream ended early or did not parse.  If modkit exited on its
            # own after a signal (killed by the OS, not by us), the input is not
            # at fault: that is retryable.  A cap hit, or a parse failure while
            # modkit is still running, is ours to stop and stays terminal.
            exited = None
            if not isinstance(exc, CallCapExceeded) and not timed_out.is_set():
                try:
                    exited = process.wait(timeout=MODKIT_EXIT_GRACE_SECONDS)
                except subprocess.TimeoutExpired:
                    exited = None
                except BaseException:
                    # Interrupted while waiting: never leave modkit running.
                    kill_process_group(process)
                    raise
            kill_process_group(process)
            if timed_out.is_set():  # a killed stream reads as a short or empty table
                raise ModkitTimedOut("modkit extraction timed out") from exc
            if exited is not None and exited < 0:
                raise ModkitKilled(f"modkit was stopped by signal {-exited}") from exc
            raise
        except BaseException as exc:
            kill_process_group(process)
            if timed_out.is_set():  # a killed stream reads as a short or empty table
                raise ModkitTimedOut("modkit extraction timed out") from exc
            raise
        finally:
            watchdog.cancel()
            if process.stdout is not None:
                process.stdout.close()
    if timed_out.is_set():
        raise ModkitTimedOut("modkit extraction timed out")
    if returncode < 0:
        raise ModkitKilled(f"modkit was stopped by signal {-returncode}")
    if returncode != 0:
        raise CellOriginPipelineError(
            "modkit extraction failed; run modkit validate on the aligned modBAM"
        )
    return normalized, counts


def _round_trips(model: Any) -> bool:
    """The model re-validates from its own JSON to an equal model (strict schema)."""

    try:
        return type(model).model_validate_json(model.model_dump_json()) == model
    except ValidationError:
        return False


def _publication_safe(payload: Any) -> bool:
    """No forbidden key (read IDs, paths, sequences) and no BAM name anywhere."""

    def keys(value: Any) -> Iterable[str]:
        if isinstance(value, dict):
            for key, item in value.items():
                yield key
                yield from keys(item)
        elif isinstance(value, list):
            for item in value:
                yield from keys(item)

    return not _FORBIDDEN_OUTPUT_KEYS.intersection(
        keys(payload)
    ) and _FORBIDDEN_OUTPUT_SUBSTRING not in json.dumps(payload, sort_keys=True)


def validation_report(
    *,
    marker_counts: Sequence[MarkerCountRow],
    classified_fragment_marker_count: int,
    atlas: AtlasUMatrix,
    deconvolution: DeconvolutionOutputV2,
    bootstrap: BootstrapResultV2 | None,
    artifacts: Sequence[tuple[DigestArtifact, Path]],
    publication: Sequence[Any] = (),
) -> ValidationReport:
    """Each publication check, computed from the result (never assumed).

    - strict_schema: every result model re-validates from its own JSON;
    - digests_verified: every input artifact still has the digest it was
      used under (re-hashed now);
    - markers_match_atlas: the fit's markers are the counted markers, each an
      atlas row, and its contributors are the atlas columns in order;
    - uxm_counts_reconciled: every marker row's U + X + M is its count, and
      the rows sum to the classified fragment-marker count;
    - fractions_normalized: fractions are in [0, 1] and sum to 1;
    - model_safe: no read ID, path, sequence or BAM name in what is published.
    """

    counted = tuple(row.marker_id for row in marker_counts)
    atlas_markers = {row.marker_id for row in atlas.rows}
    models: list[Any] = [*marker_counts, deconvolution, *publication]
    if bootstrap is not None:
        models.append(bootstrap)
    digests = True
    for artifact, path in artifacts:
        try:
            digests = digests and _sha256(path) == artifact.sha256
        except CellOriginPipelineError:
            digests = False
    fractions = [estimate.fraction for estimate in deconvolution.estimates]
    checks = {
        ValidationCheck.STRICT_SCHEMA: all(_round_trips(model) for model in models),
        ValidationCheck.DIGESTS_VERIFIED: bool(artifacts) and digests,
        ValidationCheck.MARKERS_MATCH_ATLAS: (
            bool(counted)
            and deconvolution.marker_ids == counted
            and set(counted) <= atlas_markers
            and deconvolution.atlas_id == atlas.atlas_id
            and tuple(item.cell_type_id for item in deconvolution.estimates)
            == atlas.cell_type_ids
        ),
        ValidationCheck.UXM_COUNTS_RECONCILED: (
            all(
                row.u_count + row.x_count + row.m_count == row.classified_fragment_count
                for row in marker_counts
            )
            and sum(row.classified_fragment_count for row in marker_counts)
            == classified_fragment_marker_count
        ),
        ValidationCheck.FRACTIONS_NORMALIZED: (
            all(0.0 <= value <= 1.0 for value in fractions)
            and abs(sum(fractions) - 1.0) <= 1e-9
        ),
        ValidationCheck.MODEL_SAFE: _publication_safe(
            [model.model_dump(mode="json") for model in models]
        ),
    }
    return ValidationReport(
        records=tuple(
            ValidationRecord(check=check, passed=checks[check]) for check in ValidationCheck
        )
    )


def _prefilter_notice(counts: PrefilterCounts) -> str:
    excluded = ", ".join(f"{reason} {count}" for reason, count in prefilter_summary(counts))
    return (
        f"Before modkit, {counts.records_scanned} alignment records were scanned: "
        f"{counts.written} passed and overlap a marker region, "
        f"{counts.outside_regions} passed but overlap no marker region; "
        f"excluded: {excluded}."
    )


def run_pipeline(
    config: PipelineConfig,
    *,
    fragment_hash_salt: bytes,
    software_versions: Mapping[str, str] | None = None,
    modkit: PinnedTool | None = None,
) -> CellOriginResultBundle:
    """Execute the bounded cell-origin pipeline and atomically publish JSON.

    ``modkit`` defaults to the pinned per-user toolchain (TBX-TOOL-001 when it
    is missing or not the pinned bytes).
    """

    if (
        sum(item is not None for item in (config.extract_tsv, config.aligned_modbam))
        != 1
    ):
        raise CellOriginPipelineError(
            "provide exactly one aligned modBAM or validated modkit extract TSV"
        )
    call_source = config.extract_tsv
    versions = dict(software_versions or {})
    work: Path | None = None
    prefilter: PrefilterCounts | None = None
    if config.aligned_modbam is not None:
        if config.reference_fasta is None:
            raise CellOriginPipelineError(
                "a reference FASTA is required to extract calls from an aligned modBAM"
            )
        work = _modkit_work_directory(config)
        tool = modkit if modkit is not None else resolve_modkit()
        versions["modkit"] = tool.identity.version
    resources = _load_loyfer_resources(config)

    try:
        if work is not None:
            call_source, prefilter = _extract_modbam(config, modkit=tool)
        assert call_source is not None
        calls = load_modkit_extract_calls(
            call_source,
            columns=NORMALIZED_MODKIT_COLUMNS,
            fragment_hash_salt=fragment_hash_salt,
            max_rows=config.maximum_calls,
        )
        if not calls:
            raise CellOriginPipelineError(
                "validated modkit extract contains no eligible cytosine calls"
            )
        classified = _partitioned_classification(
            calls,
            resources.markers,
            maximum_groups=config.maximum_groups,
            maximum_cpgs_per_group=config.maximum_cpgs_per_group,
        )
        if not classified.marker_counts:
            raise CellOriginPipelineError(
                "no fragments overlap a usable marker with at least four callable CpGs"
            )
        marker_ids = tuple(row.marker_id for row in classified.marker_counts)
        atlas = _filtered_atlas(resources.atlas, marker_ids)
        deconvolution = deconvolve_uxm_v2(
            classified.marker_counts,
            atlas,
            row_scale=NnlsRowScale.SQRT_COUNT,
        )
        bootstrap = bootstrap_uxm_v2(
            classified.marker_counts,
            atlas,
            deconvolution,
            replicates=config.bootstrap_replicates,
            random_seed=config.random_seed,
        )
        healthy = _load_healthy(
            config, resources.atlas.cell_type_ids, resources.labels
        )

        input_path = (
            config.extract_tsv
            if config.extract_tsv is not None
            else config.aligned_modbam
        )
        assert input_path is not None
        artifacts = [
            _artifact(input_path, "artifact.methylation-input"),
            _artifact(config.marker_bed, "artifact.loyfer-region-bed"),
            _artifact(config.marker_metadata, "artifact.loyfer-marker-metadata"),
            _artifact(config.atlas_u_matrix, "artifact.loyfer-atlas-u"),
        ]
        if config.healthy_table is not None:
            artifacts.append(
                _artifact(config.healthy_table, "artifact.healthy-table-s8")
            )

        versions.setdefault("cell-origin-pipeline", PIPELINE_VERSION)
        versions.setdefault("python", platform.python_version())
        versions.setdefault("numpy", np.__version__)
        source_ids = [config.atlas_source_id]
        if healthy is not None:
            source_ids.append(config.healthy_source_id)
        partial = classified.diagnostics.partial_input
        provenance = CellOriginProvenance(
            schema_version=SCHEMA_VERSION,
            input_artifacts=tuple(artifacts),
            source_ids=tuple(source_ids),
            software_versions=tuple(
                SoftwareVersion(
                    software_id=_identifier(name, kind="software"),
                    version=_version_token(version),
                )
                for name, version in sorted(versions.items())
            ),
            method=LOYFER_UXM_METHOD,
            uxm_thresholds=UxmThresholds(),
            input_fragment_count=len(
                {call.fragment_digest for call in calls}
            ),
            marker_overlap_count=classified.diagnostics.fragment_marker_group_count,
            classified_fragment_marker_count=(
                classified.diagnostics.classified_fragment_marker_count
            ),
            excluded_fewer_than_four_cpgs=(
                classified.diagnostics.excluded_fewer_than_four_cpgs
            ),
            partial_input=partial,
            verification_level=(
                VerificationLevel.SAMPLED_RECOMPUTED
                if partial
                else VerificationLevel.RECOMPUTED
            ),
        )
        validation = validation_report(
            marker_counts=classified.marker_counts,
            classified_fragment_marker_count=(
                classified.diagnostics.classified_fragment_marker_count
            ),
            atlas=atlas,
            deconvolution=deconvolution,
            bootstrap=bootstrap,
            artifacts=tuple(
                zip(
                    artifacts,
                    (
                        input_path,
                        config.marker_bed,
                        config.marker_metadata,
                        config.atlas_u_matrix,
                        *(() if config.healthy_table is None else (config.healthy_table,)),
                    ),
                    strict=True,
                )
            ),
            publication=(provenance,),
        )
        if not validation.passed:
            raise CellOriginPipelineError(
                "the result failed its validation checks: "
                + ", ".join(
                    record.check.value for record in validation.records if not record.passed
                )
            )
        result_digest = hashlib.sha256(
            json.dumps(
                {
                    "deconvolution": deconvolution.model_dump(mode="json"),
                    "artifacts": [
                        artifact.model_dump(mode="json") for artifact in artifacts
                    ],
                    "bootstrap": bootstrap.model_dump(mode="json"),
                },
                sort_keys=True,
                separators=(",", ":"),
            ).encode()
        ).hexdigest()
        result = CellOriginResult(
            result_id=f"cell-origin.{result_digest[:24]}",
            method=LOYFER_UXM_METHOD,
            marker_counts=classified.marker_counts,
            deconvolution=deconvolution,
            bootstrap=bootstrap,
            range_comparison=None,
            provenance=provenance,
            validation=validation,
        )
        if not _model_safe(result):
            raise CellOriginPipelineError(
                "aggregate result failed the model-safe publication check"
            )

        interval_by_id = {
            item.cell_type_id: (
                item.lower_fraction,
                item.upper_fraction,
            )
            for item in bootstrap.intervals
        }
        composition = []
        ordered_estimates = sorted(
            deconvolution.estimates,
            key=lambda item: (-item.fraction, item.cell_type_id),
        )
        for rank, estimate in enumerate(ordered_estimates, start=1):
            lower, upper = interval_by_id[estimate.cell_type_id]
            interval = next(
                item
                for item in bootstrap.intervals
                if item.cell_type_id == estimate.cell_type_id
            )
            composition.append(
                CompositionChartRow(
                    cell_type_id=estimate.cell_type_id,
                    label=resources.labels.get(
                        estimate.cell_type_id, estimate.cell_type_id
                    ),
                    rank=rank,
                    fraction=estimate.fraction,
                    percent=estimate.fraction * 100.0,
                    lower_fraction=lower,
                    upper_fraction=upper,
                    lower_percent=(
                        lower * 100.0 if lower is not None else None
                    ),
                    upper_percent=(
                        upper * 100.0 if upper is not None else None
                    ),
                    uncertainty_available=(
                        lower is not None and upper is not None
                    ),
                    uncertainty_status=interval.information_status,
                    color=PALETTE[(rank - 1) % len(PALETTE)],
                    show_by_default=rank <= DEFAULT_TOP_COMPOSITION_ROWS,
                )
            )
        estimates = {
            item.cell_type_id: item.fraction
            for item in deconvolution.estimates
        }
        healthy_rows = (
            _healthy_chart_rows(
                healthy,
                estimates,
                interval_by_id,
                resources.labels,
            )
            if healthy is not None
            else ()
        )
        notices = [
            (
                f"{resources.excluded_incomplete_count} atlas markers with one "
                "or more NA cell-type values were excluded; no imputation was used."
            ),
            (
                "Observed markers without at least four callable fragment CpGs "
                "do not enter the NNLS fit."
            ),
            (
                "Bootstrap v2 independently resamples classified fragment calls "
                "within each marker and does not preserve cross-marker molecule "
                "linkage. Unusable uncertainty is reported as unavailable."
            ),
        ]
        if prefilter is not None:
            notices.append(_prefilter_notice(prefilter))
        if resources.collapsed_duplicate_count:
            notices.append(
                f"{resources.collapsed_duplicate_count} duplicate Loyfer marker "
                "row was collapsed by genomic region after identical atlas "
                "values were verified."
            )
        method_warning = None
        range_title = (
            "Regenerated sample vs Loyfer healthy plasma donors (n=23)"
            if healthy is not None
            else None
        )
        bundle = CellOriginResultBundle(
            result=result,
            charts=ChartData(
                composition_rows=tuple(composition),
                healthy_context_rows=healthy_rows,
                composition_title="Regenerated cell-origin composition",
                range_title=range_title,
                method_warning=method_warning,
            ),
            resources=ResourceSummary(
                registered_markers=resources.registered_marker_count,
                usable_markers=len(resources.markers),
                excluded_incomplete_atlas_markers=(
                    resources.excluded_incomplete_count
                ),
                collapsed_duplicate_regions=(
                    resources.collapsed_duplicate_count
                ),
                cell_type_count=len(resources.atlas.cell_type_ids),
            ),
            notices=tuple(notices),
        )
        _atomic_write(config.output_path, bundle)
        return bundle
    except (CellOriginInputError, ValueError) as exc:
        if isinstance(exc, CellOriginPipelineError):
            raise
        raise CellOriginPipelineError(str(exc)) from exc
    finally:
        if work is not None:
            _remove_work_directory(work)


def _config_from_args(arguments: argparse.Namespace) -> PipelineConfig:
    unit = FractionUnit(arguments.healthy_unit)
    return PipelineConfig(
        marker_bed=arguments.marker_bed,
        marker_metadata=arguments.marker_metadata,
        atlas_u_matrix=arguments.atlas,
        output_path=arguments.output,
        extract_tsv=arguments.extract_tsv,
        aligned_modbam=arguments.aligned_modbam,
        healthy_table=arguments.healthy_table,
        healthy_sheet=arguments.healthy_sheet,
        healthy_fraction_unit=unit,
        maximum_calls=arguments.maximum_calls,
        maximum_groups=arguments.maximum_groups,
        bootstrap_replicates=arguments.bootstrap_replicates,
        random_seed=arguments.random_seed,
        reference_fasta=arguments.reference,
        job_directory=(
            arguments.job_dir
            if arguments.job_dir is not None
            else arguments.output.parent / "cell-origin-job"
        ),
    )


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Regenerate bounded Loyfer fragment-UXM cell-origin results."
    )
    input_group = parser.add_mutually_exclusive_group(required=True)
    input_group.add_argument("--extract-tsv", type=Path)
    input_group.add_argument("--aligned-modbam", type=Path)
    parser.add_argument(
        "--reference",
        type=Path,
        default=Path("data/local/reference/hg38.primary.fa"),
        help="FASTA the modBAM was aligned to (its .mmi is used by the alignment plan)",
    )
    parser.add_argument(
        "--job-dir",
        type=Path,
        default=None,
        help="private work directory for modkit files (default: next to --output)",
    )
    parser.add_argument(
        "--marker-bed",
        type=Path,
        default=Path("data/local/loyfer/Regions.U250.l4.hg38.bed"),
    )
    parser.add_argument(
        "--marker-metadata",
        type=Path,
        default=Path("data/local/loyfer/Markers.U250.hg38.tsv"),
    )
    parser.add_argument(
        "--atlas",
        type=Path,
        default=Path("data/local/loyfer/Atlas.U250.l4.hg38.full.tsv"),
    )
    parser.add_argument(
        "--healthy-table",
        type=Path,
        default=Path(
            "data/local/loyfer/loyfer-supplementary-tables.xlsx"
        ),
    )
    parser.add_argument("--healthy-sheet", default="Table S8")
    parser.add_argument(
        "--healthy-unit",
        choices=tuple(item.value for item in FractionUnit),
        default=FractionUnit.FRACTION.value,
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("data/local/cell-origin/result.json"),
    )
    parser.add_argument("--maximum-calls", type=int, default=DEFAULT_MAXIMUM_CALLS)
    parser.add_argument(
        "--maximum-groups", type=int, default=DEFAULT_MAXIMUM_GROUPS
    )
    parser.add_argument(
        "--bootstrap-replicates",
        type=int,
        default=DEFAULT_BOOTSTRAP_REPLICATES,
    )
    parser.add_argument("--random-seed", type=int, default=DEFAULT_RANDOM_SEED)
    parser.add_argument(
        "--print-alignment-plan",
        action="store_true",
        help="Print an argv-only plan for the configured modBAM and exit.",
    )
    return parser


def cli_main(argv: Sequence[str] | None = None) -> int:
    parser = build_argument_parser()
    arguments = parser.parse_args(argv)
    config = _config_from_args(arguments)
    if arguments.print_alignment_plan:
        if config.aligned_modbam is None:
            parser.error("--print-alignment-plan requires --aligned-modbam")
        assert config.reference_fasta is not None
        plan = build_alignment_command_plan(
            config.aligned_modbam,
            config.reference_fasta.with_name(config.reference_fasta.name + ".mmi"),
            Path("data/local/alignment-work"),
        )
        print(plan.model_dump_json(indent=2))
        return 0

    report = preflight(
        config,
        require_healthy=True,
        require_modkit=config.aligned_modbam is not None,
    )
    if not report.ready:
        print("Cell-origin regeneration is blocked:", file=sys.stderr)
        for blocker in report.blockers:
            print(f"- {blocker}", file=sys.stderr)
        return 2
    salt = os.environ.get("TRACEBACK_FRAGMENT_HASH_SALT")
    if not salt:
        print(
            "Cell-origin regeneration is blocked:\n"
            "- set TRACEBACK_FRAGMENT_HASH_SALT to a private, nonempty value",
            file=sys.stderr,
        )
        return 2
    try:
        bundle = run_pipeline(
            config,
            fragment_hash_salt=salt.encode("utf-8"),
        )
    except CellOriginPipelineError as exc:
        print(f"Cell-origin regeneration failed: {exc}", file=sys.stderr)
        return 1
    except ToolProblem as problem:
        print(
            f"Cell-origin regeneration is blocked:\n- {problem.code}: "
            f"{problem.summary}. {problem.fix}",
            file=sys.stderr,
        )
        return 2
    print(
        json.dumps(
            {
                "result_id": bundle.result.result_id,
                "output": str(config.output_path),
                "cell_types": len(bundle.charts.composition_rows),
                "classified_fragment_markers": (
                    bundle.result.provenance.classified_fragment_marker_count
                ),
            },
            indent=2,
        )
    )
    return 0


__all__ = [
    "ATLAS_METADATA_HEADER",
    "LOYFER_MAX_ROWS",
    "MARKER_METADATA_HEADER",
    "AlignmentCommandPlan",
    "AtlasU250Table",
    "CellOriginPipelineError",
    "CellOriginResultBundle",
    "ChartData",
    "CommandStep",
    "CompositionChartRow",
    "HealthyRangeChartRow",
    "PipelineConfig",
    "PreflightItem",
    "PreflightReport",
    "ResourceSummary",
    "build_alignment_command_plan",
    "build_argument_parser",
    "cli_main",
    "preflight",
    "read_atlas_u250",
    "read_marker_metadata",
    "read_marker_regions",
    "run_pipeline",
]
